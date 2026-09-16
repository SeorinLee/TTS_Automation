"""Invitation name editor parsing, numbering, API and export tests."""

import asyncio
import io
from pathlib import Path

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from gmv.api import create_app
from gmv.automation.invitation_name_editor_session import (
    InvitationNameSnapshot,
    RenameOutcome,
    TikTokInvitationNameEditorSession,
    duplicate_invitation_count,
    duplicate_invitation_numbers,
    invitation_names_from_text,
    is_unnumbered_invitation_text,
)
from gmv.config import ProfileStatus
from gmv.invitation_acceptor import (
    expand_invitation_name_edit_input,
    invitation_number_for_base,
    parse_invitation_name_edit_base,
    smallest_unused_invitation_number,
)
from gmv.store import JobStore


def test_name_edit_parser_uses_right_based_fields_and_preserves_trailing_underscore():
    parsed = parse_invitation_name_edit_base("OWNER_WITH_7MBD_0914_")
    assert parsed.owner == "OWNER_WITH"
    assert parsed.product == "7MBD"
    assert parsed.date == "0914"
    assert parsed.base_name == "OWNER_WITH_7MBD_0914_"
    assert parsed.search_query == "7MBD_0914"

    specs, errors = expand_invitation_name_edit_input(
        "D_테스트_0914_, D_테스트_0914_\nD_7MBD_0903_"
    )
    assert errors == []
    assert [item.base_name for item in specs] == ["D_테스트_0914_", "D_7MBD_0903_"]


def test_name_edit_number_matching_and_smallest_unused_are_exact():
    base = "D_7MBD_0903_"
    assert invitation_number_for_base(base, base) == 0
    assert invitation_number_for_base(base + "1", base) == 1
    assert invitation_number_for_base(base + "12", base) == 12
    assert invitation_number_for_base(base + "0", base) is None
    assert invitation_number_for_base("X_7MBD_0903_1", base) is None
    assert smallest_unused_invitation_number({1, 2, 4, 7}) == 3


def test_name_edit_reads_suffix_split_by_tiktok_markup():
    base = "D_테스트_0914_"
    assert invitation_names_from_text(base + "14", base) == (base + "14",)
    assert invitation_names_from_text(base + "\n14\nInvitation ID", base) == (base + "14",)
    assert invitation_names_from_text(base + "\u200b11", base) == (base + "11",)
    assert invitation_names_from_text(base + "\nLast modified", base) == (base,)
    assert is_unnumbered_invitation_text(base + "\nLast modified", base) is True
    assert is_unnumbered_invitation_text(base + "14\nLast modified", base) is False


def test_name_edit_finds_unnumbered_invitation_by_result_row():
    async def go():
        base = "D_테스트_0914_"
        numbered = FakeResultRow(base + "14\nLast modified")
        unnumbered = FakeResultRow(base + "\nInvitation ID")
        session = object.__new__(TikTokInvitationNameEditorSession)
        session._page = FakeResultPage([numbered, unnumbered])
        assert await session._visible_unnumbered_row(base) is unnumbered

    asyncio.run(go())


class FakeResultRow:
    def __init__(self, text, *, visible=True, cells=1):
        self.text = text
        self.visible = visible
        self.cells = cells

    async def is_visible(self):
        return self.visible

    async def inner_text(self):
        return self.text

    def locator(self, _selector):
        return FakeResultCells(self.cells)


class FakeResultCells:
    def __init__(self, count):
        self.value = count

    async def count(self):
        return self.value


class FakeResultRows:
    def __init__(self, rows):
        self.rows = rows

    async def count(self):
        return len(self.rows)

    def nth(self, index):
        return self.rows[index]


class FakeResultPage:
    def __init__(self, table_rows, role_rows=()):
        self.table_rows = FakeResultRows(table_rows)
        self.role_rows = FakeResultRows(role_rows)

    def locator(self, selector):
        if selector == "table tbody tr":
            return self.table_rows
        assert selector == "[role='rowgroup'] [role='row']"
        return self.role_rows


def test_name_edit_discovery_counts_all_visible_search_result_rows():
    async def go():
        session = object.__new__(TikTokInvitationNameEditorSession)
        session._page = FakeResultPage(
            [
                FakeResultRow("D_7MBD_0903_"),
                FakeResultRow("D_7MBD_0903_1"),
                FakeResultRow("hidden", visible=False),
            ]
        )
        assert await session._result_row_count_on_current_page(["D_7MBD_0903_"]) == 2

        session._page = FakeResultPage([], [FakeResultRow("header", cells=0)])
        assert await session._result_row_count_on_current_page(["one", "two"]) == 2

    asyncio.run(go())


def test_duplicate_invitation_suffixes_are_counted_as_remaining_repairs():
    base = "D_5CB_0915_"
    names = (
        f"{base}1",
        f"{base}2",
        f"{base}2",
        f"{base}8",
        f"{base}8",
        f"{base}8",
        base,
    )
    assert duplicate_invitation_numbers(names, base) == {2, 8}
    assert duplicate_invitation_count(names, base) == 3


def test_full_fifty_row_page_waits_for_late_pagination():
    class EmptyPagination:
        async def count(self):
            return 0

    class PaginationPage:
        def locator(self, _selector):
            return EmptyPagination()

        async def wait_for_timeout(self, _timeout):
            return None

    async def go():
        session = object.__new__(TikTokInvitationNameEditorSession)
        session._page = PaginationPage()
        session._pagination_settle_timeout_ms = 0
        attempts = 0

        async def signature():
            return "page-one"

        async def advance(_scope):
            nonlocal attempts
            attempts += 1
            return attempts == 3

        session._search_result_signature = signature
        session._advance_next = advance
        session._check_cancel = lambda: None

        assert await session._advance_scan_page(1, 50) is True
        assert attempts == 3

        attempts = 0
        assert await session._advance_scan_page(1, 35) is False
        assert attempts == 2

    asyncio.run(go())


def test_unnumbered_search_uses_stable_pagination_after_full_page():
    async def go():
        session = object.__new__(TikTokInvitationNameEditorSession)
        target = object()
        current_page = 1
        advances = []

        session._check_cancel = lambda: None

        async def emit(*_args):
            return None

        async def visible_exact(_base_name):
            return None

        async def visible_unnumbered(_base_name):
            return target if current_page == 2 else None

        async def signature():
            return f"page-{current_page}"

        async def row_count(_fallback):
            return 50 if current_page == 1 else 35

        async def advance(page_number, visible_count):
            nonlocal current_page
            advances.append((page_number, visible_count))
            current_page += 1
            return True

        session._emit = emit
        session._visible_exact_invitation = visible_exact
        session._visible_unnumbered_row = visible_unnumbered
        session._search_result_signature = signature
        session._result_row_count_on_current_page = row_count
        session._advance_scan_page = advance

        spec = parse_invitation_name_edit_base("D_5CB_0915_")
        assert await session._find_unnumbered(spec) is target
        assert advances == [(1, 50)]

    asyncio.run(go())


class FakeEditorSession:
    calls = []
    scans = []

    def __init__(self, profile):
        self.profile = profile
        self.cancel_event = None
        self.counts = {}

    async def start(self):
        return ProfileStatus.CONNECTED

    async def scan_group(self, spec, progress_cb=None):
        self.scans.append(spec.base_name)
        return InvitationNameSnapshot(
            names=(spec.base_name + "1", spec.base_name + "2", spec.base_name),
            result_total=3,
            used_numbers=frozenset({1, 2}),
            unnumbered_count=1,
            numbered_count=2,
        )

    async def rename_next(self, spec, progress_cb=None, snapshot=None, number=None):
        assert snapshot is not None
        assert number == 3
        count = self.counts.get(spec.base_name, 0)
        self.counts[spec.base_name] = count + 1
        self.calls.append((spec.base_name, self.profile.storage_root))
        if count:
            return None
        if progress_cb:
            await progress_cb("save", f"{spec.base_name} 저장", 1, 1)
        return RenameOutcome(
            original_name=spec.base_name,
            new_name=spec.base_name + "3",
            number=3,
            remaining_count=0,
            discovered_total=3,
            completed_count=3,
            snapshot=InvitationNameSnapshot(
                names=(spec.base_name + "1", spec.base_name + "2", spec.base_name + "3"),
                result_total=3,
                used_numbers=frozenset({1, 2, 3}),
                unnumbered_count=0,
                numbered_count=3,
            ),
        )

    async def close(self):
        return None


class RecoveringEditorSession(FakeEditorSession):
    attempts = 0

    async def scan_group(self, spec, progress_cb=None):
        return InvitationNameSnapshot(
            names=(spec.base_name,),
            result_total=1,
            used_numbers=frozenset(),
            unnumbered_count=1,
            numbered_count=0,
        )

    async def rename_next(self, spec, progress_cb=None, snapshot=None, number=None):
        type(self).attempts += 1
        if type(self).attempts <= 3:
            raise RuntimeError("temporary detail failure")
        return RenameOutcome(
            original_name=spec.base_name,
            new_name=spec.base_name + "1",
            number=1,
            remaining_count=0,
            discovered_total=1,
            completed_count=1,
            snapshot=InvitationNameSnapshot(
                names=(spec.base_name + "1",),
                result_total=1,
                used_numbers=frozenset({1}),
                unnumbered_count=0,
                numbered_count=1,
            ),
        )


class SequentialEditorSession(FakeEditorSession):
    assigned_numbers = []

    async def scan_group(self, spec, progress_cb=None):
        return InvitationNameSnapshot(
            names=tuple(
                [f"{spec.base_name}{number}" for number in range(1, 12)]
                + [spec.base_name] * 3
            ),
            result_total=14,
            used_numbers=frozenset(range(1, 12)),
            unnumbered_count=3,
            numbered_count=11,
        )

    async def rename_next(self, spec, progress_cb=None, snapshot=None, number=None):
        assert number is not None
        type(self).assigned_numbers.append(number)
        used = frozenset(range(1, number + 1))
        remaining = 14 - number
        return RenameOutcome(
            original_name=spec.base_name,
            new_name=f"{spec.base_name}{number}",
            number=number,
            remaining_count=remaining,
            discovered_total=14,
            completed_count=number,
            snapshot=InvitationNameSnapshot(
                names=tuple(f"{spec.base_name}{value}" for value in sorted(used))
                + tuple(spec.base_name for _ in range(remaining)),
                result_total=14,
                used_numbers=used,
                unnumbered_count=remaining,
                numbered_count=number,
            ),
        )


class LargeSequentialEditorSession(FakeEditorSession):
    assigned_numbers = []
    configured_total = 85
    configured_numbers = frozenset(range(1, 58)) - {5, 10, 20, 30, 40, 49, 52}
    configured_duplicates = frozenset({2, 8})

    def __init__(self, profile):
        super().__init__(profile)
        self.used = set(type(self).configured_numbers)
        self.duplicates = set(type(self).configured_duplicates)
        self.unnumbered = (
            type(self).configured_total - len(self.used) - len(self.duplicates)
        )

    def snapshot(self, spec):
        return InvitationNameSnapshot(
            names=tuple(f"{spec.base_name}{value}" for value in sorted(self.used))
            + tuple(f"{spec.base_name}{value}" for value in sorted(self.duplicates))
            + tuple(spec.base_name for _ in range(self.unnumbered)),
            result_total=type(self).configured_total,
            used_numbers=frozenset(self.used),
            unnumbered_count=self.unnumbered,
            numbered_count=type(self).configured_total - self.unnumbered,
        )

    async def scan_group(self, spec, progress_cb=None):
        return self.snapshot(spec)

    async def rename_next(self, spec, progress_cb=None, snapshot=None, number=None):
        assert number is not None
        assert number > max(self.used, default=0)
        type(self).assigned_numbers.append(number)
        self.used.add(number)
        if self.unnumbered:
            self.unnumbered -= 1
        else:
            self.duplicates.pop()
        remaining = self.unnumbered + len(self.duplicates)
        return RenameOutcome(
            original_name=spec.base_name,
            new_name=f"{spec.base_name}{number}",
            number=number,
            remaining_count=remaining,
            discovered_total=type(self).configured_total,
            completed_count=type(self).configured_total - remaining,
            snapshot=self.snapshot(spec),
        )


async def poll_name_edit(client, job_id, tries=400):
    for _ in range(tries):
        payload = (await client.get(f"/invitation-name-edit-jobs/{job_id}")).json()
        if payload["status"] in {"completed", "completed_with_errors", "failed", "cancelled"}:
            return payload
        await asyncio.sleep(0.01)
    raise AssertionError("name edit job did not finish")


def test_name_edit_api_runs_with_edge_profile_and_exports_verified_result(tmp_path, monkeypatch):
    async def go():
        FakeEditorSession.calls.clear()
        FakeEditorSession.scans.clear()
        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_name_editor_session_factory=FakeEditorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            parsed = await client.post(
                "/invitation-name-edit/parse",
                json={"profile_code": "US_EDGE", "invitation_text": "D_7MBD_0903_"},
            )
            assert parsed.status_code == 200
            assert parsed.json()["items"][0]["search_query"] == "7MBD_0903"

            invalid = await client.post(
                "/invitation-name-edit-jobs",
                json={"profile_code": "US_EDGE", "invitation_text": "D_7MBD_0903_1"},
            )
            assert invalid.status_code == 400

            created = await client.post(
                "/invitation-name-edit-jobs",
                json={"profile_code": "US_EDGE", "invitation_text": "D_7MBD_0903_"},
            )
            assert created.status_code == 200
            payload = await poll_name_edit(client, created.json()["job_id"])
            assert payload["status"] == "completed"
            assert payload["success"] == 1
            assert payload["invitation_name_edit_groups"][0]["completed_count"] == 3
            assert payload["invitation_name_edit_rows"][0]["new_name"] == "D_7MBD_0903_3"
            assert payload["invitation_name_edit_rows"][0]["status"] == "VERIFIED"
            edge_source = tmp_path / "profiles" / "US_EDGE"
            assert all(Path(path) == edge_source for _, path in FakeEditorSession.calls)
            assert FakeEditorSession.scans == ["D_7MBD_0903_"]
            assert len(FakeEditorSession.calls) == 1

            response = await client.get(
                f"/invitation-name-edit-jobs/{created.json()['job_id']}/download"
            )
            workbook = load_workbook(io.BytesIO(response.content))
            assert workbook.sheetnames == ["Name changes", "Summary"]
            assert workbook["Name changes"]["C2"].value == "D_7MBD_0903_3"
            assert workbook["Summary"]["E2"].value == 3

    asyncio.run(go())


def test_name_edit_job_researches_and_continues_after_one_failed_cycle(tmp_path, monkeypatch):
    async def go():
        RecoveringEditorSession.attempts = 0
        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_name_editor_session_factory=RecoveringEditorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-name-edit-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "D_7MBD_0903_"},
            )
            payload = await poll_name_edit(client, created.json()["job_id"])
            assert payload["status"] == "completed"
            assert payload["invitation_name_edit_groups"][0]["completed_count"] == 1
            assert payload["invitation_name_edit_groups"][0]["remaining_count"] == 0
            assert payload["success"] == 1
            assert RecoveringEditorSession.attempts == 4
            assert any(
                "목록 재검색 후 계속" in item["message"] for item in payload["logs"]
            )

    asyncio.run(go())


def test_name_edit_job_keeps_next_number_monotonic_until_discovered_total(tmp_path, monkeypatch):
    async def go():
        SequentialEditorSession.assigned_numbers.clear()
        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_name_editor_session_factory=SequentialEditorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-name-edit-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "D_7MBD_0903_"},
            )
            payload = await poll_name_edit(client, created.json()["job_id"])
            group = payload["invitation_name_edit_groups"][0]
            assert payload["status"] == "completed"
            assert SequentialEditorSession.assigned_numbers == [12, 13, 14]
            assert group["discovered_total"] == 14
            assert group["completed_count"] == 14
            assert group["remaining_count"] == 0

    asyncio.run(go())


@pytest.mark.parametrize(
    ("total", "used_numbers", "duplicates", "expected_first", "expected_last"),
    [
        (
            85,
            frozenset(range(1, 58)) - {5, 10, 20, 30, 40, 49, 52},
            frozenset({2, 8}),
            58,
            92,
        ),
        (125, frozenset(range(1, 58)), frozenset(), 58, 125),
    ],
)
def test_name_edit_large_groups_continue_after_highest_number_without_filling_gaps(
    tmp_path,
    monkeypatch,
    total,
    used_numbers,
    duplicates,
    expected_first,
    expected_last,
):
    async def go():
        LargeSequentialEditorSession.assigned_numbers.clear()
        LargeSequentialEditorSession.configured_total = total
        LargeSequentialEditorSession.configured_numbers = used_numbers
        LargeSequentialEditorSession.configured_duplicates = duplicates
        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_name_editor_session_factory=LargeSequentialEditorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-name-edit-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "D_5CB_0915_"},
            )
            payload = await poll_name_edit(client, created.json()["job_id"], tries=1_000)
            group = payload["invitation_name_edit_groups"][0]
            assigned = LargeSequentialEditorSession.assigned_numbers
            assert payload["status"] == "completed"
            assert assigned == list(range(expected_first, expected_last + 1))
            assert len(assigned) == total - len(used_numbers)
            assert len(assigned) == len(set(assigned))
            assert not (set(assigned) & set(used_numbers))
            assert group["discovered_total"] == total
            assert group["completed_count"] == total
            assert group["remaining_count"] == 0

    asyncio.run(go())
