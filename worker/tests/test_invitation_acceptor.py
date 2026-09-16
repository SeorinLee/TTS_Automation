"""Invitation Acceptor parser, safety, persistence and control tests."""

import asyncio
import io
from pathlib import Path

from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from gmv.api import create_app
from gmv.automation.invitation_inspector_session import (
    INVITATION_SEARCH_TABS,
    CreatorDetails,
    InvitationInspection,
    InvitationNotFoundError,
    TikTokInvitationInspectorSession,
    creator_activity_flags,
    creator_count_warning,
    creator_identity_lines,
    inspect_json_for_creators,
    invitation_suffix_query,
    invitation_tab_pattern,
    is_invitation_suffix_match,
    merge_creator_pages,
    target_url_from_authenticated_page,
)
from gmv.config import ProfileStatus, get_profile
from gmv.invitation_acceptor import (
    expand_invitation_input,
    group_orders_by_product,
    is_exact_invitation_match,
    parse_invitation_name,
)
from gmv.invitation_acceptor_service import creator_is_member
from gmv.store import JobStore


def test_right_based_parser_range_order_and_product_grouping():
    parsed = parse_invitation_name("OWNER_WITH_UNDERSCORE_SZP_0810_15")
    assert parsed.owner == "OWNER_WITH_UNDERSCORE"
    assert (parsed.product, parsed.date, parsed.number) == ("SZP", "0810", "15")
    specs, errors = expand_invitation_input(
        "PJH_SZP_0810_1~3, D_MC_0720_15\nD_KTB_0430_1~2"
    )
    assert errors == []
    assert [spec.full_name for spec in specs] == [
        "PJH_SZP_0810_1",
        "PJH_SZP_0810_2",
        "PJH_SZP_0810_3",
        "D_MC_0720_15",
        "D_KTB_0430_1",
        "D_KTB_0430_2",
    ]
    assert group_orders_by_product(specs) == {
        "szp": [1, 2, 3],
        "mc": [4],
        "ktb": [5, 6],
    }


def test_exact_match_never_accepts_numeric_suffix_or_owner_variants():
    target = "PJH_SZP_0810_1"
    assert is_exact_invitation_match("pjh_szp_0810_1", target)
    assert not is_exact_invitation_match("PJH_SZP_0810_10", target)
    assert not is_exact_invitation_match("PJH_SZP_0810_11", target)
    assert not is_exact_invitation_match("ABC_SZP_0810_1", target)


def test_invitation_suffix_query_respects_tiktok_ten_character_limit():
    assert INVITATION_SEARCH_TABS == ("Ongoing",)
    assert invitation_suffix_query("PJH_SZPEUKA_0814_59") == "A_0814_59"
    assert invitation_suffix_query("PJH_SZPEUKA_0814_55") == "A_0814_55"
    assert invitation_suffix_query("SHORT") == "SHORT"
    assert invitation_tab_pattern("Ongoing").fullmatch("Ongoing 1")
    assert invitation_tab_pattern("Ongoing").fullmatch("Ongoing")
    assert not invitation_tab_pattern("Ongoing").fullmatch("Completed 1")


def test_suffix_match_accepts_rendered_invitation_name_but_not_longer_number():
    assert is_invitation_suffix_match("D_KTB_0707_22", "B_0707_22")
    assert is_invitation_suffix_match("D_KT B_0707_22\nOngoing", "B_0707_22")
    assert not is_invitation_suffix_match("D_KTB_0707_220", "B_0707_22")


def test_suffix_matched_result_is_opened_without_full_name_exact_locator():
    class MissingRow:
        async def count(self):
            return 0

    class SuffixNode:
        clicked = False

        async def inner_text(self):
            return "D_KT B_0707_22"

        def locator(self, _selector):
            return MissingRow()

        async def click(self):
            self.clicked = True

    class FakePage:
        async def wait_for_timeout(self, _milliseconds):
            return None

    async def go():
        session = TikTokInvitationInspectorSession(get_profile("US_CHROME"))
        session._page = FakePage()
        node = SuffixNode()

        await session._open_candidate(node, "D_KTB_0707_22")

        assert node.clicked

    asyncio.run(go())


def test_edge_detail_tab_is_adopted_and_original_list_tab_is_remembered():
    class FakePage:
        def __init__(self, url):
            self.url = url
            self.events = []
            self.front = False

        def is_closed(self):
            return False

        def on(self, event, _callback):
            self.events.append(event)

        async def bring_to_front(self):
            self.front = True

        async def wait_for_load_state(self, *_args, **_kwargs):
            return None

        async def wait_for_timeout(self, _milliseconds):
            return None

    class FakeContext:
        def __init__(self, pages):
            self.pages = pages

    async def go():
        list_page = FakePage(
            "https://affiliate-us.tiktok.com/affiliate/collaboration/target-invitation"
        )
        detail_page = FakePage(
            "https://affiliate-us.tiktok.com/affiliate/collaboration/target-invitation/detail"
        )
        session = TikTokInvitationInspectorSession(get_profile("US_EDGE"))
        session._page = list_page
        session._context = FakeContext([list_page, detail_page])

        adopted = await session._adopt_opened_detail_page([list_page], list_page)

        assert adopted is True
        assert session._page is detail_page
        assert session._return_page is list_page
        assert detail_page.front is True
        assert detail_page.events == ["response"]

    asyncio.run(go())


def test_suffix_search_presses_enter_waits_for_fresh_result_and_returns_exact_invitation():
    class FakeNode:
        def __init__(self, text="", present=False):
            self.text = text
            self.present = present

        async def count(self):
            return 1 if self.present else 0

        def nth(self, _index):
            return self

        async def is_visible(self):
            return self.present

        async def inner_text(self):
            return self.text

    class FakeBody:
        def __init__(self, page):
            self.page = page

        async def inner_text(self):
            return self.page.body_text

    class FakePage:
        def __init__(self, requested, suffix):
            self.requested = requested
            self.suffix = suffix
            self.body_text = "Ongoing invitation list"
            self.suffix_pending = False
            self.suffix_ready = False
            self.pending_ticks = 0

        def locator(self, selector):
            assert selector == "body"
            return FakeBody(self)

        def get_by_text(self, text, exact=False):
            assert exact is True
            present = self.suffix_ready and text == self.requested
            return FakeNode(self.requested, present)

        async def wait_for_load_state(self, *_args, **_kwargs):
            return None

        async def wait_for_timeout(self, _milliseconds):
            if not self.suffix_pending or self.suffix_ready:
                return
            self.pending_ticks += 1
            if self.pending_ticks >= 4:
                self.suffix_ready = True
                self.body_text = f"suffix results {self.requested}"

    class FakeSearch:
        def __init__(self, page):
            self.page = page
            self.value = ""
            self.presses = []

        async def fill(self, value):
            self.value = value

        async def input_value(self):
            return self.value

        async def press(self, key):
            self.presses.append((self.value, key))
            if self.value == self.page.suffix:
                # TikTok keeps the broad result visible while this request is in flight.
                self.page.suffix_pending = True
            else:
                self.page.body_text = f"broad results for {self.value}"

    class TestSession(TikTokInvitationInspectorSession):
        async def _select_invitation_tab(self, _tab_name):
            return True

        async def _find_search(self):
            return self.search

        async def _advance_next(self, _scope):
            return False

    async def go():
        spec = parse_invitation_name("OWNER_KTB_0804_31")
        suffix = invitation_suffix_query(spec.full_name)
        session = TestSession(get_profile("US_CHROME"))
        session._page = FakePage(spec.full_name, suffix)
        session.search = FakeSearch(session._page)

        node = await session._find_exact_invitation(spec, None)

        assert await node.inner_text() == spec.full_name
        assert session.search.presses == [("KTB", "Enter"), ("B_0804_31", "Enter")]
        assert session._page.pending_ticks >= 4

    asyncio.run(go())


def test_target_route_preserves_authenticated_shop_parameters():
    url = target_url_from_authenticated_page(
        get_profile("US_EDGE"),
        "https://seller-us.tiktok.com/affiliate/landing?shop_region=US",
        "https://affiliate-us.tiktok.com/affiliate/collaboration/target-invitation?shop_region=US&shop_id=7495830785034323995&route_migration=1&tab=1",
    )
    assert url == (
        "https://affiliate-us.tiktok.com/affiliate/collaboration/target-invitation"
        "?shop_region=US&shop_id=7495830785034323995&route_migration=1&tab=1"
    )


def test_creator_pages_preserve_invited_order_and_dedupe_by_id():
    merged = merge_creator_pages(
        [
            [CreatorDetails("creator_a", "A", "10000001", "US")],
            [
                CreatorDetails("renamed_a", "A2", "10000001", "US"),
                CreatorDetails("creator_b", "B", "10000002", "US"),
            ],
        ]
    )
    assert [item.creator for item in merged] == ["creator_a", "creator_b"]


def test_sample_request_pages_collect_all_creators_and_dedupe_by_id():
    class FakePage:
        def __init__(self):
            self.current = 0

        async def wait_for_timeout(self, _milliseconds):
            return None

    class SampleSession(TikTokInvitationInspectorSession):
        pages = [
            [CreatorDetails("creator_a", "A", "10000001", "US")],
            [
                CreatorDetails("renamed_a", "A2", "10000001", "US"),
                CreatorDetails("creator_b", "B", "10000002", "US"),
            ],
        ]

        async def _creator_rows(self, scope):
            return self.pages[scope.current]

        async def _advance_next(self, scope):
            if scope.current + 1 >= len(self.pages):
                return False
            scope.current += 1
            return True

    async def go():
        session = SampleSession(get_profile("US_CHROME"))
        page = FakePage()
        session._page = page
        progress = []

        async def report(phase, message, page_number=0, _total=0):
            progress.append((phase, message, page_number))

        creators = await session._collect_creator_pages(
            page,
            "Sample Requests",
            report,
            phase="samples",
        )

        assert [item.creator for item in creators] == ["creator_a", "creator_b"]
        assert [item[0] for item in progress] == ["samples", "samples"]

    asyncio.run(go())


def test_creator_count_mismatch_is_a_non_fatal_warning():
    assert creator_count_warning(49, 48) == "경고: 카드 49 / 실제 48"
    assert creator_count_warning(48, 48) is None
    assert creator_count_warning(49, 0) is None
    assert creator_count_warning(None, 48) is None


def test_creator_network_json_extracts_id_nickname_and_region():
    creators = inspect_json_for_creators(
        {"data": {"list": [{"unique_id": "milannlynn", "nickname": "Milan", "creator_id": "74954847", "region": "US"}]}}
    )
    assert creators == [CreatorDetails("milannlynn", "Milan", "74954847", "US")]


def test_invited_creator_activity_columns_become_exact_o_flags():
    assert creator_activity_flags(["1 product", "0 products", "View details"]) == (True, False)
    assert creator_activity_flags(["2 products", "3 videos"]) == (True, True)
    assert creator_activity_flags(["0 products", "0 contents"]) == (False, False)


def test_creator_identity_skips_sample_expand_cell_and_timing_marker():
    assert creator_identity_lines(
        ["", ".\nhealthsharebear\nTheShareBear\nPPS: 4.2/5.0", "80"]
    )[:2] == ["healthsharebear", "TheShareBear"]


def test_unique_nickname_fallback_matches_when_tab_hides_creator_id():
    invited = [
        CreatorDetails("creator_a", "Unique Nick", "74950001", "US"),
        CreatorDetails("creator_b", "Other", "74950002", "US"),
    ]
    added = [CreatorDetails("masked_name", "Unique Nick", "", "US")]
    assert creator_is_member(invited[0], added, invited, set())
    assert not creator_is_member(invited[1], added, invited, set())


class FakeAcceptorSession:
    calls = []

    def __init__(self, profile):
        self.profile = profile
        self.cancel_event = None

    async def start(self):
        return ProfileStatus.CONNECTED

    async def inspect_invitation(self, spec, progress_cb=None):
        type(self).calls.append((spec.full_name, self.profile.storage_root))
        if progress_cb:
            await progress_cb("search", f"{spec.product} search", 1, 2)
        creators = [
            CreatorDetails("creator_a", "Creator A", "74950001", "US"),
            CreatorDetails("creator_b", "Creator B", "74950002", "US"),
        ]
        return InvitationInspection(
            creators=creators,
            added_creators=[creators[0]],
            posted_creators=[creators[1]],
            sample_creators=[creators[0]],
            added_product_keys={creators[0].identity},
            posted_content_keys={creators[1].identity},
            sample_request_keys={creators[0].identity},
        )

    async def close(self):
        return None


async def poll(client, job_id, terminal, tries=400):
    for _ in range(tries):
        payload = (await client.get(f"/invitation-accept-jobs/{job_id}")).json()
        if payload["status"] in terminal:
            return payload
        await asyncio.sleep(0.01)
    raise AssertionError(f"job did not reach {terminal}")


def test_acceptor_api_persists_order_uses_edge_source_and_exports_two_sheets(tmp_path, monkeypatch):
    async def go():
        FakeAcceptorSession.calls.clear()
        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_acceptor_session_factory=FakeAcceptorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            parsed = await client.post(
                "/invitation-accept/parse",
                json={"profile_code": "US_EDGE", "invitation_text": "PJH_SZP_0810_1~2"},
            )
            assert [item["number"] for item in parsed.json()["items"]] == ["1", "2"]
            created = await client.post(
                "/invitation-accept-jobs",
                json={"profile_code": "US_EDGE", "invitation_text": "PJH_SZP_0810_1~2"},
            )
            assert created.status_code == 200
            payload = await poll(client, created.json()["job_id"], {"completed"})
            assert payload["concurrency"] == 1
            assert [state["status"] for state in payload["invitation_accept_states"]] == [
                "SUCCESS",
                "SUCCESS",
            ]
            assert len(payload["invitation_creator_rows"]) == 4
            assert payload["invitation_accept_states"][0]["sample_requests_count"] == 1
            assert payload["invitation_creator_rows"][0]["sample_sent"] is True
            assert payload["invitation_creator_rows"][1]["sample_sent"] is False
            assert [call[0] for call in FakeAcceptorSession.calls] == [
                "PJH_SZP_0810_1",
                "PJH_SZP_0810_2",
            ]
            edge_source = tmp_path / "profiles" / "US_EDGE"
            assert all(Path(call[1]) == edge_source for call in FakeAcceptorSession.calls)
            assert not (tmp_path / "runtime").exists()
            response = await client.get(
                f"/invitation-accept-jobs/{created.json()['job_id']}/download"
            )
            workbook = load_workbook(io.BytesIO(response.content))
            assert workbook.sheetnames == ["Results", "Errors"]
            assert [workbook["Results"].cell(row, 1).value for row in (2, 3)] == ["SZP", "SZP"]
            assert workbook["Results"]["G2"].value == "O"
            assert workbook["Results"]["H3"].value == "O"
            assert workbook["Results"]["I1"].value == "Sample sent"
            assert workbook["Results"]["I2"].value == "O"
            assert workbook["Results"]["I3"].value is None
            assert workbook["Errors"].max_row == 1

    asyncio.run(go())


def test_invalid_input_is_rejected_before_browser_start(tmp_path):
    async def go():
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_acceptor_session_factory=FakeAcceptorSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            response = await client.post(
                "/invitation-accept-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "BAD_NAME"},
            )
            assert response.status_code == 400
            assert response.json()["code"] == "INVALID_INVITATION_INPUT"

    asyncio.run(go())


def test_pause_then_resume_continues_next_invitation(tmp_path, monkeypatch):
    async def go():
        entered = asyncio.Event()
        release = asyncio.Event()

        class GateSession(FakeAcceptorSession):
            async def inspect_invitation(self, spec, progress_cb=None):
                if spec.number == "1":
                    entered.set()
                    await release.wait()
                return await super().inspect_invitation(spec, progress_cb)

        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_acceptor_session_factory=GateSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-accept-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "PJH_SZP_0810_1~2"},
            )
            job_id = created.json()["job_id"]
            await asyncio.wait_for(entered.wait(), timeout=3)
            assert (await client.post(f"/invitation-accept-jobs/{job_id}/pause")).status_code == 200
            release.set()
            for _ in range(200):
                paused = (await client.get(f"/invitation-accept-jobs/{job_id}")).json()
                if paused["status"] == "paused" and paused["processed"] == 1:
                    break
                await asyncio.sleep(0.01)
            else:
                raise AssertionError("job did not pause between invitations")
            assert (await client.post(f"/invitation-accept-jobs/{job_id}/resume")).status_code == 200
            completed = await poll(client, job_id, {"completed"})
            assert completed["processed"] == 2

    asyncio.run(go())


def test_cancel_keeps_checkpoint_and_marks_remaining_items(tmp_path, monkeypatch):
    async def go():
        entered = asyncio.Event()
        release = asyncio.Event()

        class GateSession(FakeAcceptorSession):
            async def inspect_invitation(self, spec, progress_cb=None):
                entered.set()
                await release.wait()
                return await super().inspect_invitation(spec, progress_cb)

        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_acceptor_session_factory=GateSession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-accept-jobs",
                json={"profile_code": "US_CHROME", "invitation_text": "PJH_SZP_0810_1~2"},
            )
            job_id = created.json()["job_id"]
            await asyncio.wait_for(entered.wait(), timeout=3)
            assert (await client.post(f"/invitation-accept-jobs/{job_id}/cancel")).status_code == 200
            release.set()
            cancelled = await poll(client, job_id, {"cancelled"})
            assert cancelled["processed"] == 2
            assert [item["status"] for item in cancelled["invitation_accept_states"]] == [
                "SUCCESS",
                "CANCELLED",
            ]
            assert (await client.get(f"/invitation-accept-jobs/{job_id}/download")).status_code == 200

    asyncio.run(go())


def test_retry_runs_only_failed_items(tmp_path, monkeypatch):
    async def go():
        calls = []

        class FlakySession(FakeAcceptorSession):
            async def inspect_invitation(self, spec, progress_cb=None):
                calls.append(spec.full_name)
                if spec.number == "1" and calls.count(spec.full_name) == 1:
                    raise InvitationNotFoundError("not found on first run")
                return await super().inspect_invitation(spec, progress_cb)

        monkeypatch.setenv("GMV_PROFILE_ROOT", str(tmp_path / "profiles"))
        monkeypatch.setenv("GMV_RUNTIME_PROFILE_ROOT", str(tmp_path / "runtime"))
        app = create_app(
            JobStore(str(tmp_path / "jobs")),
            invitation_acceptor_session_factory=FlakySession,
        )
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            created = await client.post(
                "/invitation-accept-jobs",
                json={"profile_code": "US_EDGE", "invitation_text": "PJH_SZP_0810_1~2"},
            )
            job_id = created.json()["job_id"]
            first = await poll(client, job_id, {"completed_with_errors"})
            assert [item["status"] for item in first["invitation_accept_states"]] == [
                "NOT_FOUND",
                "SUCCESS",
            ]
            assert (await client.post(f"/invitation-accept-jobs/{job_id}/retry")).status_code == 200
            final = await poll(client, job_id, {"completed"})
            assert [item["status"] for item in final["invitation_accept_states"]] == [
                "SUCCESS",
                "SUCCESS",
            ]
            assert calls == ["PJH_SZP_0810_1", "PJH_SZP_0810_2", "PJH_SZP_0810_1"]

    asyncio.run(go())
