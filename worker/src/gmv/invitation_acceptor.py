"""Invitation Acceptor parsing, safe identity matching and XLSX export."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_INVITATION_RE = re.compile(r"^(.*?)_([^_]+)_([^_]+)_(\d+)$")
_INVITATION_EDIT_RE = re.compile(r"^(.*?)_([^_]+)_([^_]+)_$")
_RANGE_RE = re.compile(r"^(.*?)_(\d+)~(\d+)$")


class InvitationInputError(ValueError):
    def __init__(self, detail: str, code: str = "INVALID_INVITATION_INPUT"):
        super().__init__(detail)
        self.detail = detail
        self.code = code


@dataclass(frozen=True)
class InvitationSpec:
    order: int
    full_name: str
    owner: str
    product: str
    date: str
    number: str


@dataclass(frozen=True)
class InvitationNameEditSpec:
    """One unnumbered invitation-name family, parsed with the lookup's right-based fields."""

    order: int
    base_name: str
    owner: str
    product: str
    date: str

    @property
    def search_query(self) -> str:
        return f"{self.product}_{self.date}"


def normalize_invitation_name(value: object) -> str:
    """Normalize only whitespace and case; punctuation and numeric suffix stay significant."""
    return str(value or "").strip().casefold()


def parse_invitation_name(full_name: str, *, order: int = 1) -> InvitationSpec:
    name = str(full_name or "").strip()
    match = _INVITATION_RE.fullmatch(name)
    if not match or not match.group(1):
        raise InvitationInputError(
            f"형식을 확인해주세요: {name or '(빈 값)'}. 예: PJH_SZP_0810_1"
        )
    return InvitationSpec(
        order=order,
        full_name=name,
        owner=match.group(1),
        product=match.group(2),
        date=match.group(3),
        number=match.group(4),
    )


def parse_invitation_name_edit_base(
    full_name: str,
    *,
    order: int = 1,
) -> InvitationNameEditSpec:
    """Parse ``Owner_Product_Date_`` using the same right-based layout as lookup names."""
    name = str(full_name or "").strip()
    match = _INVITATION_EDIT_RE.fullmatch(name)
    if not match or not match.group(1):
        raise InvitationInputError(
            f"형식을 확인해주세요: {name or '(빈 값)'}. 예: D_테스트_0914_"
        )
    if len(name) + 1 > 30:
        raise InvitationInputError(
            f"번호를 붙이면 TikTok 30자 제한을 넘습니다: {name}"
        )
    return InvitationNameEditSpec(
        order=order,
        base_name=name,
        owner=match.group(1),
        product=match.group(2),
        date=match.group(3),
    )


def expand_invitation_name_edit_input(
    raw: str,
) -> tuple[list[InvitationNameEditSpec], list[str]]:
    """Parse newline/comma separated edit bases and preserve first-appearance order."""
    tokens = [part.strip() for part in re.split(r"[\r\n,]+", str(raw or "")) if part.strip()]
    specs: list[InvitationNameEditSpec] = []
    errors: list[str] = []
    seen: set[str] = set()
    for token in tokens:
        key = normalize_invitation_name(token)
        if key in seen:
            continue
        try:
            spec = parse_invitation_name_edit_base(token, order=len(specs) + 1)
        except InvitationInputError as exc:
            errors.append(exc.detail)
            continue
        seen.add(key)
        specs.append(spec)
    return specs, errors


def invitation_number_for_base(actual: str, base_name: str) -> int | None:
    """Return 0 for the unnumbered base, a positive suffix number, or ``None``."""
    actual_key = normalize_invitation_name(actual)
    base_key = normalize_invitation_name(base_name)
    if actual_key == base_key:
        return 0
    match = re.fullmatch(re.escape(base_key) + r"(\d+)", actual_key)
    if not match:
        return None
    number = int(match.group(1))
    return number if number > 0 else None


def smallest_unused_invitation_number(used: set[int]) -> int:
    """Choose the smallest positive suffix absent from the latest live search result."""
    candidate = 1
    while candidate in used:
        candidate += 1
    return candidate


def write_invitation_name_edit_results(groups, rows, output_path: str) -> str:
    """Write a continuously usable checkpoint workbook for invitation rename jobs."""
    workbook = Workbook()
    results = workbook.active
    results.title = "Name changes"
    results.append(
        ["Group", "Original invitation", "New invitation", "Number", "Status", "Result"]
    )
    for row in rows:
        results.append(
            [
                row.base_name,
                row.original_name,
                row.new_name,
                row.number,
                row.status,
                row.message,
            ]
        )

    summary = workbook.create_sheet("Summary")
    summary.append(
        ["No", "Input", "Search", "Found", "Completed", "Remaining", "Failed", "Status", "Message"]
    )
    for group in sorted(groups, key=lambda item: item.order):
        summary.append(
            [
                group.order,
                group.base_name,
                group.search_query,
                group.discovered_total,
                group.completed_count,
                group.remaining_count,
                group.failed_count,
                group.status,
                group.message,
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="355B6E")
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 10), 48)
    workbook.save(output_path)
    return output_path


def expand_invitation_input(raw: str) -> tuple[list[InvitationSpec], list[str]]:
    """Expand final-number ranges and preserve the user's first-appearance order."""
    tokens = [part.strip() for part in re.split(r"[\r\n,]+", str(raw or "")) if part.strip()]
    expanded: list[str] = []
    errors: list[str] = []
    for token in tokens:
        range_match = _RANGE_RE.fullmatch(token)
        if range_match:
            start, end = int(range_match.group(2)), int(range_match.group(3))
            if start > end or end - start > 5000:
                errors.append(f"올바르지 않은 범위: {token}")
                continue
            expanded.extend(f"{range_match.group(1)}_{number}" for number in range(start, end + 1))
        else:
            expanded.append(token)

    specs: list[InvitationSpec] = []
    seen: set[str] = set()
    for full_name in expanded:
        key = normalize_invitation_name(full_name)
        if key in seen:
            continue
        try:
            spec = parse_invitation_name(full_name, order=len(specs) + 1)
        except InvitationInputError as exc:
            errors.append(exc.detail)
            continue
        seen.add(key)
        specs.append(spec)
    return specs, errors


def is_exact_invitation_match(actual: str, requested: str) -> bool:
    return normalize_invitation_name(actual) == normalize_invitation_name(requested)


def group_orders_by_product(specs: list[InvitationSpec]) -> dict[str, list[int]]:
    grouped: dict[str, list[int]] = {}
    for spec in specs:
        grouped.setdefault(spec.product.casefold(), []).append(spec.order)
    return grouped


def write_invitation_accept_results(states, output_path: str, creator_rows=None) -> str:
    workbook = Workbook()
    results = workbook.active
    results.title = "Results"
    results.append(
        [
            "Keyword",
            "Invitation",
            "Creator",
            "Nickname",
            "Creator ID",
            "Region",
            "Added products",
            "Posted content",
            "Sample sent",
        ]
    )
    for row in sorted(creator_rows or [], key=lambda item: item.order):
        results.append(
            [
                row.keyword,
                row.invitation_name,
                row.creator,
                row.nickname,
                row.creator_id,
                row.region or row.market,
                "O" if row.added_products else None,
                "O" if row.posted_content else None,
                "O" if row.sample_sent else None,
            ]
        )

    errors = workbook.create_sheet("Errors")
    errors.append(["Order", "Invitation", "Error", "Time"])
    successful = {"SUCCESS"}
    for state in sorted(states, key=lambda item: item.order):
        if state.status in successful | {"QUEUED", "PROCESSING"}:
            continue
        errors.append([state.order, state.invitation_name, state.message or state.status, state.processed_at])

    for sheet in (results, errors):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4968ED")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 16,
        "B": 32,
        "C": 24,
        "D": 24,
        "E": 20,
        "F": 12,
        "G": 18,
        "H": 18,
        "I": 16,
    }.items():
        results.column_dimensions[column].width = width
    for column, width in {"A": 10, "B": 32, "C": 60, "D": 24}.items():
        errors.column_dimensions[column].width = width
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path
